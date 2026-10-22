"""Build the editable paper using the supplied PDF's academic layout.

Run from the repository root. A final Word pagination/TOC refresh is performed
by render_word_paper.ps1, using the installed Microsoft Word application.
"""
from pathlib import Path
import re
import subprocess
import tempfile
import sys
import os
import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt, Mm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SUBMISSION = '--submission' in sys.argv
args = [a for a in sys.argv[1:] if a != '--submission']
IS_REPORT = bool(args)
SOURCE = Path(args[0]).resolve() if args else ROOT / 'paper/AR_EAURP_Research_Paper.md'
OUT = Path(args[1]).resolve() if len(args) > 1 else ROOT / 'AR_EAURP_Research_Paper_Revised.docx'
ROMAN = 'Times New Roman'
BLACK = '000000'


def child(parent, tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn('w:' + k), str(v))
    parent.append(e)
    return e


def font(obj, family=ROMAN, size=None, bold=None):
    obj.font.name = family
    if size is not None:
        obj.font.size = Pt(size)
    if bold is not None:
        obj.font.bold = bold
    rpr = obj._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = child(rpr, 'w:rFonts')
    for attr in list(rf.attrib):
        del rf.attrib[attr]
    for name in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rf.set(qn('w:' + name), family)
    c = rpr.find(qn('w:color'))
    if c is None:
        c = child(rpr, 'w:color')
    c.attrib.clear()
    c.set(qn('w:val'), BLACK)


def configure_style(d, name, size=10, family=ROMAN, bold=False,
                    before=0, after=6, line=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    if name not in d.styles:
        d.styles.add_style(name, 1)
    s = d.styles[name]
    font(s, family, size, bold)
    if not hasattr(s, 'paragraph_format'):
        return s
    f = s.paragraph_format
    f.space_before = Pt(before)
    f.space_after = Pt(after)
    f.line_spacing = Pt(line)
    f.alignment = align
    f.keep_with_next = False
    f.keep_together = False
    f.widow_control = True
    for b in s._element.xpath('.//w:pBdr'):
        b.getparent().remove(b)
    return s


def field(p, instruction, text=''):
    run = p.add_run()
    child(run._r, 'w:fldChar', fldCharType='begin')
    ins = child(p.add_run()._r, 'w:instrText')
    ins.set(qn('xml:space'), 'preserve')
    ins.text = instruction
    child(p.add_run()._r, 'w:fldChar', fldCharType='separate')
    p.add_run(text)
    child(p.add_run()._r, 'w:fldChar', fldCharType='end')


source = SOURCE.read_text(encoding='utf-8-sig')
source = source.replace('\ufeff', '')
# The Markdown starts at H2; Word's first two heading levels should be 12/10 pt.
source = re.sub(r'^(#{2,6}) ', lambda m: m[1][1:] + ' ', source, flags=re.M)
with tempfile.TemporaryDirectory(prefix='ar-paper-') as tmp:
    md = Path(tmp) / 'paper.md'
    md.write_text(source, encoding='utf-8')
    subprocess.run([
        pypandoc.get_pandoc_path(), str(md), '-f',
        'markdown+yaml_metadata_block+pipe_tables+tex_math_dollars',
        '-t', 'docx', '--standalone',
        '--resource-path=' + os.pathsep.join([str(ROOT / 'paper'), str(ROOT)]),
        '-o', str(OUT)], check=True, cwd=ROOT / 'paper')

d = Document(OUT)
sec = d.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Pt(57.6)
sec.footer_distance = Pt(25.8)
sec.header_distance = Pt(20)

for s in d.styles:
    if s.type in (1, 2, 3):
        font(s)
        for b in s._element.xpath('.//w:pBdr'):
            b.getparent().remove(b)
for name in ('Normal', 'Body Text', 'First Paragraph', 'Compact'):
    configure_style(d, name)
configure_style(d, 'Title', 17.22, 'Times New Roman', before=33, after=19,
                line=21.92, align=WD_ALIGN_PARAGRAPH.CENTER)
configure_style(d, 'Author', 11.96, 'Times New Roman', after=15,
                line=14.35, align=WD_ALIGN_PARAGRAPH.CENTER)
configure_style(d, 'Date', 11.96, 'Times New Roman', after=22,
                line=14.35, align=WD_ALIGN_PARAGRAPH.CENTER)
if SUBMISSION:
    configure_style(d, 'Title', 17, before=0, after=12, line=20,
                    align=WD_ALIGN_PARAGRAPH.CENTER)
    configure_style(d, 'Author', 10, after=12, line=12,
                    align=WD_ALIGN_PARAGRAPH.CENTER)
for level, size, family, before in [(1, 11.96, 'Times New Roman', 17.9),
                                  (2, 9.96, ROMAN, 11.9),
                                  (3, 9.96, ROMAN, 9)]:
    s = configure_style(d, f'Heading {level}', size, family, True,
                        before, 6, size * 1.2, WD_ALIGN_PARAGRAPH.LEFT)
    s.paragraph_format.keep_with_next = True
for n in ('Source Code', 'Verbatim Char'):
    configure_style(d, n, 9, 'Courier New', line=11,
                    align=WD_ALIGN_PARAGRAPH.LEFT)
configure_style(d, 'Caption', 9.96, line=12, after=10,
                align=WD_ALIGN_PARAGRAPH.CENTER)
configure_style(d, 'ContentsTitle', 14.35, 'Times New Roman', True,
                before=0, after=6, line=17.2, align=WD_ALIGN_PARAGRAPH.LEFT)
for lev in range(1, 4):
    s = configure_style(d, f'TOC {lev}', 9.96, before=0, after=0,
                        line=11.95, align=WD_ALIGN_PARAGRAPH.LEFT)
    s.paragraph_format.left_indent = Pt(0 if lev == 1 else 22.4 * (lev - 1))
    s.paragraph_format.first_line_indent = Pt(0)
    tabs = s.paragraph_format.tab_stops
    tabs.clear_all()
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    tabs.add_tab_stop(Pt(480.08), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

for p in d.paragraphs:
    fmt = p.paragraph_format
    is_rule = any(e.get('{urn:schemas-microsoft-com:office:office}hr') == 't'
                  for e in p._p.iter('{urn:schemas-microsoft-com:vml}rect'))
    # Remove Pandoc/Word theme color, Arial title and accidental justification.
    if p.style.name == 'Title' and not IS_REPORT:
        p.clear()
        p.add_run('AR-EAURP: An Adversarial-Resilient and Energy-Harvesting-Aware')
        p.add_run().add_break()
        p.add_run('Routing Stack for Mobile Ad Hoc Networks')
    elif p.text.strip() == 'AR-EAURP Project':
        p.style = d.styles['Author']
    elif p.text.strip() == '22 October 2026':
        p.style = d.styles['Date']
    elif p.text.startswith('Manuscript type:'):
        fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif p.style.name == 'Source Code':
        fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif p._p.xpath('.//w:drawing'):
        fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt.space_before, fmt.space_after = Pt(10), Pt(4)
        fmt.line_spacing = 1.0
        fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
        fmt.keep_with_next = True
    elif p._p.xpath('.//m:oMathPara'):
        fmt.line_spacing = 1.0
        fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
        fmt.space_before, fmt.space_after = Pt(6), Pt(6)
    elif p.text.startswith('Table ') and len(p.text) < 240:
        p.style = d.styles['Caption']
        fmt.keep_with_next = True
        fmt.space_before, fmt.space_after = Pt(8), Pt(4)
    for r in p.runs:
        if p.style.name == 'Title':
            font(r, 'Times New Roman', 17.22, False)
        elif p.style.name in ('Author', 'Date'):
            font(r, 'Times New Roman', 11.96, False)
        elif p.style.name.startswith('Heading'):
            font(r, 'Times New Roman' if p.style.name == 'Heading 1' else ROMAN,
                 11.96 if p.style.name == 'Heading 1' else 9.96, True)
        elif r.style and r.style.name in ('Verbatim Char', 'Source Code'):
            font(r, 'Courier New', 9)
        else:
            font(r, ROMAN)
    for border in p._p.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
    if is_rule:
        p.clear()
        fmt.left_indent = fmt.right_indent = Pt(120)
        fmt.space_before, fmt.space_after, fmt.line_spacing = Pt(6), Pt(12), Pt(1)
        borders = child(p._p.get_or_add_pPr(), 'w:pBdr')
        child(borders, 'w:bottom', val='single', sz=4, color='808080')

# Compact, centered, journal-style rules around tables; no shaded or vertical grid.
widths = [[100, 165, 215], [80, 260, 70, 70], [110, 120, 110, 140],
          [130, 155, 155], [150, 110, 110, 110], [130, 87.5, 87.5, 87.5, 87.5],
          [110, 180, 190]]
for index, table in enumerate(d.tables):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if IS_REPORT:
        n = len(table.columns)
        first = 160 if n > 2 else 240
        if table.cell(0, 0).text.startswith('Hostile'):
            first = 65
        ws = [first] + [(480 - first) / (n - 1)] * (n - 1) if n > 1 else [480]
        if n == 4 and table.cell(0, 1).text.startswith('Factors'):
            ws = [85, 255, 70, 70]
        if n == 6 and table.cell(0, 0).text == 'A minus B':
            ws = [190, 60, 110, 45, 35, 40]
    else:
        ws = widths[index]
    props = table._tbl.tblPr
    for node in list(props):
        if node.tag in (qn('w:tblBorders'), qn('w:tblCellMar'), qn('w:tblStyle')):
            props.remove(node)
    border = child(props, 'w:tblBorders')
    for side in ('top', 'bottom'):
        child(border, 'w:' + side, val='single', sz='5', color='000000')
    for side in ('left', 'right', 'insideH', 'insideV'):
        child(border, 'w:' + side, val='nil')
    margins = child(props, 'w:tblCellMar')
    for side, val in [('top', 45), ('bottom', 45), ('left', 60), ('right', 60)]:
        child(margins, 'w:' + side, w=val, type='dxa')
    for col, w in zip(table.columns, ws):
        col.width = Pt(w)
    for ri, row in enumerate(table.rows):
        tr = row._tr.get_or_add_trPr()
        child(tr, 'w:cantSplit')
        if ri == 0:
            child(tr, 'w:tblHeader')
        for ci, cell in enumerate(row.cells):
            cell.width = Pt(ws[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            for shade in cp.findall(qn('w:shd')):
                cp.remove(shade)
            if ri == 0:
                cb = child(cp, 'w:tcBorders')
                child(cb, 'w:bottom', val='single', sz='4', color='000000')
            for p in cell.paragraphs:
                f = p.paragraph_format
                f.alignment = WD_ALIGN_PARAGRAPH.LEFT
                f.space_before, f.space_after, f.line_spacing = Pt(0), Pt(0), Pt(11)
                keep_table = IS_REPORT and table.cell(0, 0).text == 'A minus B'
                f.keep_with_next = ri < len(table.rows) - 1 if keep_table else ri == 0
                for r in p.runs:
                    font(r, ROMAN, 8 if len(ws) >= 7 else 9, ri == 0)

for shape in d.inline_shapes:
    max_w, max_h = Pt(480), Pt(305)
    scale = min(1, max_w / shape.width, max_h / shape.height)
    shape.width = int(shape.width * scale)
    shape.height = int(shape.height * scale)

# Research manuscripts begin with their abstract; the companion report has contents.
if not SUBMISSION:
    start = next(p for p in d.paragraphs if p.style.name == 'Heading 1')
    toc_title = start.insert_paragraph_before('Contents', 'ContentsTitle')
    toc = start.insert_paragraph_before()
    toc.style = d.styles['Normal']
    toc.paragraph_format.space_after = Pt(0)
    field(toc, ' TOC \\o "1-2" \\h \\z \\u ')
    pagebreak = start.insert_paragraph_before()
    pagebreak.add_run().add_break(WD_BREAK.PAGE)
    pagebreak.paragraph_format.space_after = Pt(0)
    pagebreak.paragraph_format.line_spacing = Pt(1)

footer = sec.footer.paragraphs[0]
footer.clear()
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field(footer, ' PAGE ', '1')
for r in footer.runs:
    font(r, ROMAN, 9.96)
settings = d.settings.element
for tag, val in [('updateFields', 'true'), ('autoHyphenation', 'true'),
                 ('embedTrueTypeFonts', 'true'), ('saveSubsetFonts', 'false')]:
    e = settings.find(qn('w:' + tag))
    if e is None:
        e = child(settings, 'w:' + tag)
    e.set(qn('w:val'), val)
d.save(OUT)
print(OUT)
